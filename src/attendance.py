import os
from datetime import datetime
from openpyxl import Workbook, load_workbook

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ATTENDANCE_DIR = os.path.join(BASE_DIR, "..", "Data", "attendance")
FILE_PATH = os.path.join(ATTENDANCE_DIR, "attendance.xlsx")


def mark_attendance(name):
    os.makedirs(ATTENDANCE_DIR, exist_ok=True)
    if not os.path.exists(FILE_PATH):
        wb = Workbook()
        ws = wb.active
        ws.title = "attendance"
        ws.append(["Name", "In_Time", "Out_Time"])
        wb.save(FILE_PATH)

    wb = load_workbook(FILE_PATH)
    ws = wb.active

    time_now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    found = False

    for row in ws.iter_rows(min_row=2):
        if row[0].value == name:
            found = True

            if row[2].value is None:
                row[2].value = time_now
                wb.save(FILE_PATH)
                print(f"[Out] {name} at {time_now}")
            else:
                print(f"{name} already marked IN & OUT")

            return

    if not found:
        ws.append([name, time_now, None])
        wb.save(FILE_PATH)
        print(f"[In] {name} at {time_now}")
